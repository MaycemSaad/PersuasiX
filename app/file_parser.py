"""File parsing utilities for PersuasiX — PDF, DOCX, TXT extraction."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from loguru import logger


@dataclass
class ParsedDocument:
    """Container for a parsed document."""
    filename: str
    text: str
    page_count: int
    word_count: int
    file_type: str
    success: bool
    error: str = ""


def parse_file(file_path: str) -> ParsedDocument:
    """Parse a file and extract its text content."""
    path = Path(file_path)
    if not path.exists():
        return ParsedDocument(
            filename=path.name, text="", page_count=0, word_count=0,
            file_type="", success=False, error="File not found.",
        )

    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return _parse_pdf(path)
        elif ext == ".docx":
            return _parse_docx(path)
        elif ext in (".txt", ".text", ".md", ".csv"):
            return _parse_text(path)
        elif ext == ".html" or ext == ".htm":
            return _parse_html(path)
        else:
            return ParsedDocument(
                filename=path.name, text="", page_count=0, word_count=0,
                file_type=ext, success=False,
                error=f"Unsupported file format: {ext}. Supported: PDF, DOCX, TXT, HTML.",
            )
    except Exception as e:
        logger.error(f"Failed to parse {path.name}: {e}")
        return ParsedDocument(
            filename=path.name, text="", page_count=0, word_count=0,
            file_type=ext, success=False, error=str(e),
        )


def _parse_pdf(path: Path) -> ParsedDocument:
    """Extract text from a PDF file."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return ParsedDocument(
            filename=path.name, text="", page_count=0, word_count=0,
            file_type=".pdf", success=False,
            error="pypdf not installed. Run: pip install pypdf",
        )

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text.strip())

    full_text = "\n\n".join(pages)
    if not full_text.strip():
        return ParsedDocument(
            filename=path.name, text="", page_count=len(reader.pages),
            word_count=0, file_type=".pdf", success=False,
            error="Could not extract text from PDF. It may be scanned/image-based.",
        )

    logger.info(f"Parsed PDF: {path.name} — {len(reader.pages)} pages, {len(full_text.split())} words")
    return ParsedDocument(
        filename=path.name, text=full_text, page_count=len(reader.pages),
        word_count=len(full_text.split()), file_type=".pdf", success=True,
    )


def _parse_docx(path: Path) -> ParsedDocument:
    """Extract text from a DOCX file."""
    try:
        from docx import Document
    except ImportError:
        return ParsedDocument(
            filename=path.name, text="", page_count=0, word_count=0,
            file_type=".docx", success=False,
            error="python-docx not installed. Run: pip install python-docx",
        )

    doc = Document(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n\n".join(paragraphs)

    if not full_text.strip():
        return ParsedDocument(
            filename=path.name, text="", page_count=0, word_count=0,
            file_type=".docx", success=False,
            error="DOCX file appears to be empty.",
        )

    logger.info(f"Parsed DOCX: {path.name} — {len(paragraphs)} paragraphs, {len(full_text.split())} words")
    return ParsedDocument(
        filename=path.name, text=full_text, page_count=0,
        word_count=len(full_text.split()), file_type=".docx", success=True,
    )


def _parse_text(path: Path) -> ParsedDocument:
    """Read a plain text file."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = path.read_text(encoding=encoding)
            if text.strip():
                logger.info(f"Parsed TXT: {path.name} — {len(text.split())} words")
                return ParsedDocument(
                    filename=path.name, text=text.strip(), page_count=0,
                    word_count=len(text.split()), file_type=path.suffix.lower(),
                    success=True,
                )
        except (UnicodeDecodeError, UnicodeError):
            continue

    return ParsedDocument(
        filename=path.name, text="", page_count=0, word_count=0,
        file_type=path.suffix.lower(), success=False,
        error="Could not decode the text file.",
    )


def _parse_html(path: Path) -> ParsedDocument:
    """Extract text from an HTML file."""
    import re
    raw = path.read_text(encoding="utf-8", errors="replace")
    text = re.sub(r"<[^>]+>", " ", raw)
    text = re.sub(r"\s+", " ", text).strip()

    if not text:
        return ParsedDocument(
            filename=path.name, text="", page_count=0, word_count=0,
            file_type=".html", success=False, error="HTML file appears empty.",
        )

    return ParsedDocument(
        filename=path.name, text=text, page_count=0,
        word_count=len(text.split()), file_type=".html", success=True,
    )
